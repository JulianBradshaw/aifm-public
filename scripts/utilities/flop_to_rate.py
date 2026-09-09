#!/usr/bin/env python3
"""Reconstruct an instantaneous training *rate* series from a cumulative training-FLOP series.

We have a time series F(t) -- e.g. "FLOP used in the largest AI training run to date" --
specified on an integer-decimal-year grid (2017.0, 2018.0, ...). In this codebase,
exponentially-growing quantities are interpolated *in log space* between grid points:

    F is the target cumulative; we instead want a RATE r(t) [FLOP/s] such that the model's
    log-space interpolation of r, integrated over time, reproduces F on the grid.

Key identity
------------
If the rate is log-linearly interpolated between knots r_i, r_{i+1} over [t_i, t_{i+1}]:

    r(t_i + u*dt) = r_i^{1-u} * r_{i+1}^{u},   u in [0, 1]

then the FLOP accumulated over that interval is

    int_{t_i}^{t_{i+1}} r(t) * S dt = S * dt_i * L(r_i, r_{i+1})

where S is seconds-per-year and L(a, b) = (b - a) / (ln b - ln a) is the *logarithmic mean*
(with L(a, a) = a). So the exact grid-matching condition for each interval is

    L(r_i, r_{i+1}) = A_i := (F_{i+1} - F_i) / (S * dt_i).

Given r_0 this is a forward recursion (each step has a unique positive r_{i+1}), so one
boundary convention determines the whole rate series and reproduces F exactly at every
grid point. The interval-average of a log-linear rate being the log-mean of its endpoints
is the whole trick: it gives a *continuous* rate series (unlike differentiating the
log-linear F, which is exact but discontinuous at knots) that still integrates exactly.

Boundary conventions (``--method``)
-----------------------------------
``first-segment`` (default): r_0 = g_0 * F_0 / S with g_0 = (ln F_1 - ln F_0) / (t_1 - t_0),
    i.e. the derivative implied by extrapolating the first log-linear segment. Equivalent to
    "assume exponential growth at the 2017->2018 rate for all past time and integrate from
    -infinity"; recovers the ordinary derivative exactly when F is a pure exponential. This
    is GPT-5.5's proposal.

``smoothest``: the solution family is 1-DOF (parameterised by r_0) and the forward recursion
    is exact for *any* r_0, so we instead pick the r_0 that minimises the roughness of log r
    (sum of squared second differences). This removes the boundary-seeded ringing the
    forward recursion otherwise exhibits while keeping the integral match exact.

Either way the cumulative offset before the first grid point is taken to be C(t_0) = F_0
(the fictitious continuous run has already accumulated F_0 by the first year).
"""

from __future__ import annotations

import argparse
import csv
import os
import sys

import numpy as np
from scipy.optimize import brentq, minimize_scalar

# Repo root (scripts/utilities/<this> -> repo root) for importing model_config.
_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if _REPO_ROOT not in sys.path:
    sys.path.insert(0, _REPO_ROOT)

from progress_model import config as cfg  # noqa: E402

# Model convention: FLOP_PER_H100E_YEAR = (FLOP/s per H100e) * (seconds / year).
SECONDS_PER_YEAR = 3.15576e7  # matches the model's year convention (model_config uses 3.15576e7)


def _log_logmean(log_a: float, log_b: float) -> float:
    """Return log(L(exp(log_a), exp(log_b))) for the logarithmic mean L, computed stably."""
    d = log_b - log_a
    if abs(d) < 1e-8:
        # L -> geometric mean as the endpoints coincide; second-order correction.
        return 0.5 * (log_a + log_b) + d * d / 24.0
    hi = max(log_a, log_b)
    lo = min(log_a, log_b)
    # log((e^hi - e^lo) / |d|) = hi + log1p(-e^(lo-hi)) - log|d|
    return hi + np.log1p(-np.exp(lo - hi)) - np.log(abs(d))


def _solve_next_log_rate(log_r_i: float, log_A_i: float) -> float:
    """Solve log L(r_i, r_{i+1}) = log_A_i for log r_{i+1} (unique positive root)."""

    def f(log_r_next: float) -> float:
        return _log_logmean(log_r_i, log_r_next) - log_A_i

    lo, hi = log_r_i - 300.0, log_r_i + 300.0
    # L is monotone increasing in r_{i+1}; widen the bracket until it straddles the root.
    while f(lo) > 0:
        lo -= 300.0
    while f(hi) < 0:
        hi += 300.0
    return brentq(f, lo, hi, xtol=1e-13, rtol=1e-15)


def _forward_log_rates(log_r0: float, log_A: np.ndarray) -> np.ndarray:
    """Forward recursion: given the initial log-rate and per-interval log-averages, return all log-rates."""
    log_r = [float(log_r0)]
    for la in log_A:
        log_r.append(_solve_next_log_rate(log_r[-1], la))
    return np.array(log_r)


def _interval_log_averages(years: np.ndarray, log_F: np.ndarray, seconds_per_year: float):
    """Per-interval log(A_i) where A_i = (F_{i+1} - F_i) / (S * dt_i), plus the year deltas."""
    dt = np.diff(years)
    # log(F_{i+1} - F_i), stable for large magnitudes.
    log_dF = log_F[1:] + np.log1p(-np.exp(log_F[:-1] - log_F[1:]))
    log_A = log_dF - np.log(seconds_per_year) - np.log(dt)
    return log_A, dt


def infer_log_rate_knots(
    years: np.ndarray,
    log_F: np.ndarray,
    seconds_per_year: float = SECONDS_PER_YEAR,
    method: str = "first-segment",
) -> np.ndarray:
    """Infer log(rate) [rate in FLOP/s] at each grid point so log-interp + integrate reproduces F.

    Args:
        years: strictly increasing decimal-year grid.
        log_F: natural log of the cumulative target FLOP at each grid point (strictly increasing).
        seconds_per_year: year convention used to convert the per-year integral into FLOP.
        method: "first-segment" (GPT-5.5 boundary) or "smoothest" (min-roughness r_0).
    """
    years = np.asarray(years, dtype=float)
    log_F = np.asarray(log_F, dtype=float)

    if years.ndim != 1 or years.shape != log_F.shape:
        raise ValueError("years and log_F must be 1-D arrays of the same length")
    if np.any(np.diff(years) <= 0):
        raise ValueError("years must be strictly increasing")
    bad = np.where(np.diff(log_F) <= 0)[0]
    if bad.size:
        rows = ", ".join(f"{years[i]:.2f}->{years[i + 1]:.2f}" for i in bad[:8])
        more = "" if bad.size <= 8 else f" (+{bad.size - 8} more)"
        raise ValueError(
            "cumulative FLOP must be strictly increasing to be expressible as the integral of a "
            f"positive rate; non-increasing segment(s): {rows}{more}"
        )

    log_A, dt = _interval_log_averages(years, log_F, seconds_per_year)

    # GPT-5.5 boundary: r_0 = g_0 * F_0 / S, g_0 = first log-linear segment slope (per year).
    g0 = (log_F[1] - log_F[0]) / dt[0]
    log_r0_first_segment = np.log(g0) + log_F[0] - np.log(seconds_per_year)

    if method == "first-segment":
        return _forward_log_rates(log_r0_first_segment, log_A)

    if method == "smoothest":
        def roughness(log_r0: float) -> float:
            lr = _forward_log_rates(log_r0, log_A)
            d2 = np.diff(lr, 2)
            return float(np.dot(d2, d2))

        # The min-roughness r_0 sits near the first-segment value; search a wide log-window.
        res = minimize_scalar(
            roughness,
            bounds=(log_r0_first_segment - 15.0, log_r0_first_segment + 15.0),
            method="bounded",
            options={"xatol": 1e-10},
        )
        return _forward_log_rates(res.x, log_A)

    raise ValueError(f"unknown method {method!r} (expected 'first-segment' or 'smoothest')")


def roundtrip_cumulative(
    years: np.ndarray, rate: np.ndarray, f0: float, seconds_per_year: float = SECONDS_PER_YEAR
) -> np.ndarray:
    """Integrate the log-interpolated rate on the grid, returning the reconstructed cumulative FLOP."""
    years = np.asarray(years, dtype=float)
    rate = np.asarray(rate, dtype=float)
    dt = np.diff(years)
    a, b = rate[:-1], rate[1:]
    log_a, log_b = np.log(a), np.log(b)
    # Logarithmic mean, with the geometric-mean limit when endpoints coincide.
    logmean = np.where(np.abs(log_b - log_a) < 1e-9, np.sqrt(a * b), (b - a) / (log_b - log_a))
    seg = seconds_per_year * dt * logmean
    return np.concatenate([[f0], f0 + np.cumsum(seg)])


# ---------------------------------------------------------------------------
# Input loading
# ---------------------------------------------------------------------------

def _to_flop(values: np.ndarray, unit: str) -> np.ndarray:
    """Convert an input column to absolute FLOP."""
    values = np.asarray(values, dtype=float)
    if unit == "flop":
        return values
    if unit == "log10-flop":
        return np.power(10.0, values)
    if unit == "h100e":
        return values * cfg.FLOP_PER_H100E_YEAR
    raise ValueError(f"unknown unit {unit!r} (expected flop | log10-flop | h100e)")


def load_series(path: str, time_col: str, value_col: str, unit: str, sheet: str | None):
    """Load (years, F_flop) from a CSV or .xlsx file."""
    if path.lower().endswith(".xlsx"):
        try:
            import openpyxl  # noqa: WPS433 (optional dependency)
        except ImportError as exc:  # pragma: no cover
            raise SystemExit("reading .xlsx requires openpyxl (pip install openpyxl)") from exc
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        header = [str(c).strip() if c is not None else "" for c in rows[0]]
        ti, vi = header.index(time_col), header.index(value_col)
        years, vals = [], []
        for r in rows[1:]:
            if r[ti] is None or r[vi] is None:
                continue
            try:
                y = float(r[ti])
            except (TypeError, ValueError):
                continue
            years.append(y)
            vals.append(float(r[vi]))
        years, vals = np.array(years), np.array(vals)
    else:
        with open(path, newline="") as fh:
            reader = csv.DictReader(fh)
            years, vals = [], []
            for row in reader:
                years.append(float(row[time_col]))
                vals.append(float(row[value_col]))
        years, vals = np.array(years), np.array(vals)

    order = np.argsort(years)
    return years[order], _to_flop(vals[order], unit)


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    # The xlsx forecast this used to default to is no longer in the repo (see
    # docs/input-series-provenance.md); pass the series explicitly.
    p.add_argument("--input", required=True, help="input CSV or .xlsx of the cumulative-FLOP series")
    p.add_argument("--sheet", default="Export format", help="sheet name for .xlsx input")
    p.add_argument("--time-col", default="time", help="time column name")
    p.add_argument("--value-col", default="training_compute", help="cumulative-FLOP column name")
    p.add_argument(
        "--unit",
        default="log10-flop",
        choices=["flop", "log10-flop", "h100e"],
        help="unit of the value column (xlsx Export format training_compute is log10-flop)",
    )
    p.add_argument(
        "--method",
        default="first-segment",
        choices=["first-segment", "smoothest"],
        help="boundary convention for r_0 (default: GPT-5.5's first-segment)",
    )
    p.add_argument("--seconds-per-year", type=float, default=SECONDS_PER_YEAR)
    p.add_argument("--output", help="write a CSV with columns: year, flop, flop_per_s, h100e")
    p.add_argument("--assert-tol", type=float, default=1e-9, help="max round-trip rel error to tolerate")
    args = p.parse_args(argv)

    years, flop = load_series(args.input, args.time_col, args.value_col, args.unit, args.sheet)
    log_rate = infer_log_rate_knots(years, np.log(flop), args.seconds_per_year, args.method)
    rate = np.exp(log_rate)  # FLOP/s

    # FLOP/s per H100e = FLOP_PER_H100E_YEAR / seconds_per_year (the per-second throughput of one H100e).
    flop_per_s_per_h100e = cfg.FLOP_PER_H100E_YEAR / args.seconds_per_year
    h100e = rate / flop_per_s_per_h100e

    recon = roundtrip_cumulative(years, rate, float(flop[0]), args.seconds_per_year)
    rel = np.abs(recon - flop) / flop
    max_rel = float(rel.max())

    print(f"loaded {len(years)} points {years[0]:.2f}..{years[-1]:.2f} from {os.path.basename(args.input)}")
    print(f"method={args.method}  round-trip max rel error on grid = {max_rel:.2e}")
    for yy in (years[0], 2025.0, 2030.0, 2040.0, years[-1]):
        i = int(np.argmin(np.abs(years - yy)))
        print(f"  {years[i]:7.2f}: rate={rate[i]:.3e} FLOP/s = {h100e[i]:.3e} H100e")

    if args.output:
        with open(args.output, "w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["year", "flop", "flop_per_s", "h100e"])
            for y, f, r, he in zip(years, flop, rate, h100e):
                w.writerow([f"{y:.4f}", repr(float(f)), repr(float(r)), repr(float(he))])
        print(f"wrote {args.output}")

    if max_rel > args.assert_tol:
        print(f"ERROR: round-trip error {max_rel:.2e} exceeds tolerance {args.assert_tol:.0e}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
